"""Assemble the A1-portrait single-page poster (PPTX canvas 59.4 x 84.1 cm).
Layout uses an explicit y-cursor per column; leftover column space is distributed
evenly over section gaps so all columns end at an even height above the footer."""
import math
from pathlib import Path
from PIL import Image
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(r"C:\MOV\metal_defect_vlm\output")
FIGS = OUT / "poster_figs"
LOGO = OUT / "logos" / "uob_crest.png"
NAVY = RGBColor(0x17, 0x36, 0x5D)
GREY = RGBColor(0x44, 0x44, 0x44)
BODY = RGBColor(0x1A, 0x1A, 0x1A)

PW, PH = 59.4, 84.1
COLW = 18.6
X1, X2, X3 = 1.2, 20.4, 39.6

prs = Presentation()
prs.slide_width = Cm(PW)
prs.slide_height = Cm(PH)
slide = prs.slides.add_slide(prs.slide_layouts[6])

CM_PER_PT = 0.03528
CHAR_W = 0.50
LINE_H = 1.18


def est_h(text, width_cm, size):
    cpl = max(10, int(width_cm / (size * CHAR_W * CM_PER_PT)))
    lines = 0
    for seg in text.split("\n"):
        lines += max(1, math.ceil(len(seg) / cpl))
    return lines * size * LINE_H * CM_PER_PT


class Col:
    def __init__(self, x, y0, w=COLW):
        self.x, self.y, self.w = x, y0, w
        self.items = []      # (shape, y0)
        self.gaps = []       # y positions of section starts

    def rec(self, shape, y0):
        self.items.append((shape, y0))

    def tb(self, h):
        box = slide.shapes.add_textbox(Cm(self.x), Cm(self.y), Cm(self.w), Cm(h))
        self.rec(box, self.y)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Cm(0.05)
        tf.margin_top = tf.margin_bottom = Cm(0.0)
        return tf

    def text(self, txt, size=14.5, bold=False, italic=False, color=BODY, after=0.14,
             align=PP_ALIGN.LEFT):
        h = est_h(txt, self.w, size) + after
        tf = self.tb(h)
        p = tf.paragraphs[0]
        p.alignment = align
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = "Calibri"
        self.y += h
        return h

    def bullets(self, items, size=14.5, after=0.08):
        for it in items:
            self.text("- " + it, size=size, after=after)

    def head(self, txt, size=22, before=0.4):
        self.y += before
        self.gaps.append(self.y)
        self.text(txt, size=size, bold=True, color=NAVY, after=0.12)

    def pic(self, path, w, cap=None, cap_size=13, center=True):
        im = Image.open(path)
        h = w * im.size[1] / im.size[0]
        x = self.x + (self.w - w) / 2 if center else self.x
        shp = slide.shapes.add_picture(str(path), Cm(x), Cm(self.y), width=Cm(w))
        self.rec(shp, self.y)
        self.y += h + 0.08
        if cap:
            self.text(cap, size=cap_size, italic=True, color=GREY, after=0.2)
        return h

    def shade(self, txt, fill, size=14, pad=0.3, bold_prefix=None):
        h = est_h(txt, self.w - 0.7, size) + pad
        sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(self.x), Cm(self.y),
                                    Cm(self.w), Cm(h))
        self.rec(sp, self.y)
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
        sp.line.fill.background()
        sp.shadow.inherit = False
        tf = sp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Cm(0.35)
        tf.margin_top = tf.margin_bottom = Cm(0.15)
        p = tf.paragraphs[0]
        p.space_after = Pt(0)
        if bold_prefix:
            r0 = p.add_run()
            r0.text = bold_prefix
            r0.font.size = Pt(size)
            r0.font.bold = True
            r0.font.name = "Calibri"
            r0.font.color.rgb = BODY
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.name = "Calibri"
        r.font.color.rgb = BODY
        self.y += h + 0.18
        return h

    def distribute(self, target):
        """Spread leftover space evenly over section gaps so the column ends at target."""
        delta = target - self.y
        if delta <= 0.3 or not self.gaps:
            return
        n = len(self.gaps)
        for shape, y0 in self.items:
            k = sum(1 for g in self.gaps if g <= y0 + 1e-6)
            shape.top = Cm(y0 + delta * k / n)
        self.y = target


def table_img(rows, fname, widths, fontsize=13.5, wrap=None):
    """Render a table image; rows wrap to multiple lines with per-row heights."""
    import textwrap
    if wrap:
        rows = [[textwrap.fill(cell, wrap[j]) if isinstance(cell, str) else cell
                 for j, cell in enumerate(row)] for row in rows]
    rowlines = [max(str(c).count("\n") + 1 for c in row) for row in rows]
    total = sum(rowlines)
    fig, ax = plt.subplots(figsize=(8.6, 0.40 * total + 0.35))
    ax.axis("off")
    t = ax.table(cellText=rows, cellLoc="center", colWidths=widths, loc="center")
    t.auto_set_font_size(False)
    t.set_fontsize(fontsize)
    base = 0.40 / (0.40 * total + 0.35)
    for i, nl in enumerate(rowlines):
        for j in range(len(rows[0])):
            t[i, j].set_height(base * nl * 1.25)
    for j in range(len(rows[0])):
        c = t[0, j]
        c.set_facecolor("#17365D")
        c.set_text_props(color="white", weight="bold")
    for i in range(1, len(rows)):
        for j in range(len(rows[0])):
            t[i, j].set_facecolor("#F7F9FC" if i % 2 else "#FFFFFF")
    fig.savefig(FIGS / fname, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return FIGS / fname


tbl_counts = table_img([
    ["Class", "Dataset", "Test"], ["Punching hole", "219", "23"], ["Welding line", "273", "28"],
    ["Crescent gap", "226", "24"], ["Water spot", "289", "30"], ["Oil spot", "204", "21"],
    ["Silk spot", "651", "66"], ["Inclusion", "216", "23"], ["Rolled pit", "31", "4"],
    ["Crease", "53", "6"], ["Waist folding", "150", "15"], ["Total", "2,312", "240"]],
    "tbl_counts.png", [0.5, 0.25, 0.25], fontsize=13)
tbl_overall = table_img([
    ["Model", "Acc.", "Ma-P", "Ma-R", "Ma-F1", "W-F1", "Err."],
    ["Qwen2.5-VL-7B", "86.67", "86.19", "83.05", "84.14", "86.63", "32"],
    ["Gemma3-4B", "87.08", "85.99", "84.00", "83.59", "86.75", "31"],
    ["InternVL3-8B", "86.67", "85.57", "84.44", "84.80", "86.52", "32"]],
    "tbl_overall.png", [0.32, 0.12, 0.12, 0.12, 0.12, 0.12, 0.08], fontsize=13.5)
tbl_arch = table_img([
    ["Model", "Architecture rationale"],
    ["Qwen2.5-VL-7B", "Dynamic-resolution ViT + M-RoPE; fine-grained visual features"],
    ["Gemma3-4B", "SigLIP encoder, compressed visual tokens; resource-conscious route"],
    ["InternVL3-8B", "Native multimodal pre-training; dynamic high-res tiling"],
    ["LLaVA-OneVision-7B", "SigLIP + MLP projector + AnyRes; excluded (integration)"]],
    "tbl_arch.png", [0.26, 0.74], fontsize=12.5, wrap=[15, 56])
tbl_deploy = table_img([
    ["Model", "Environment", "Latency", "Observation"],
    ["Qwen2.5-VL-7B (7B)", "Local RTX 5070 Ti", "1.690 s/img", "Locally executable"],
    ["Gemma3-4B (4B)", "Local RTX 5070 Ti", "2.529 s/img", "Locally executable"],
    ["InternVL3-8B (8B)", "BlueBEAR A100", "62.723 s/img", "HPC after local load failure"]],
    "tbl_deploy.png", [0.26, 0.24, 0.17, 0.33], fontsize=12.5, wrap=[20, 18, 12, 26])
tbl_keys = table_img([
    ["87.08%", "Highest accuracy - Gemma3-4B"],
    ["84.80%", "Highest Macro-F1 - InternVL3-8B"],
    ["86.19%", "Highest Macro precision - Qwen2.5-VL-7B"],
    ["0.41 pp", "Maximum accuracy gap (one test image)"],
    ["69.57%", "Inclusion recall, all three models"]],
    "tbl_keys.png", [0.2, 0.8], fontsize=14)

# ================= header =================
im = Image.open(LOGO)
slide.shapes.add_picture(str(LOGO), Cm(1.4), Cm(1.0), width=Cm(4.4))
box = slide.shapes.add_textbox(Cm(7.0), Cm(0.9), Cm(51.2), Cm(5.4))
tf = box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.space_after = Pt(2)
r = p.add_run()
r.text = "Comparative Evaluation of Vision-Language Models for Metal Surface Defect Classification"
r.font.size = Pt(34)
r.font.bold = True
r.font.color.rgb = NAVY
r.font.name = "Calibri"
for txt, size, style, color in [
    ("Expert Prompting and 4-bit QLoRA Adaptation of Three Architecture-Distinct VLMs on the GC10-DET Dataset",
     18, "i", GREY),
    ("Shixiang Wang, Yongjing Wang, Richard Fox", 18, "b", BODY),
    ("School of Electronic, Electrical and Systems Engineering (EESE), University of Birmingham", 15, "n", GREY)]:
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(2)
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(size)
    r.font.name = "Calibri"
    r.font.color.rgb = color
    r.font.bold = style == "b"
    r.font.italic = style == "i"

colq = Col(1.2, 6.6, w=57.0)
colq.shade("Under a shared dataset split, Expert Prompting strategy and 4-bit QLoRA framework, how do "
           "three architecture-distinct open-source VLMs compare in predictive performance, class-wise behaviour and "
           "deployment feasibility?", "DCE6F1", size=17, bold_prefix="Research question:  ")

banner = Col(4.0, colq.y + 0.2, w=51.4)
banner.pic(FIGS / "fig2_pipeline.png", 51.4,
           cap="Figure 1. Experimental workflow. All three VLMs were adapted with 4-bit QLoRA in LLaMA-Factory on a shared "
               "stratified split and evaluated on the identical 240-image test set; each model retained its native vision "
               "processor, so adaptation settings remained partly model-specific.", cap_size=13)
YBODY = banner.y + 0.2

# ---- footer strip geometry (references span all columns at the bottom) ----
STRIP_H = 4.6
STRIP_TOP = PH - 1.0 - STRIP_H
TARGET = STRIP_TOP - 0.5
cols = []

# ================= column 1 =================
c = Col(X1, YBODY)
c.head("1. Background", before=0)
c.text("Automated metal-surface inspection is essential for industrial quality control, yet conventional vision models "
       "need task-specific architectures for a fixed class set. Vision-language models (VLMs) offer a unified framework "
       "for image understanding and instruction-based classification, but general benchmarks do not reflect specialised "
       "industrial imagery: metal defects show low contrast, subtle texture change and strong inter-class similarity, "
       "combined with substantial class imbalance.")
c.text("Three open-source VLMs were selected along the architectural spectrum rather than by popularity, so the comparison "
       "covers distinct design strategies under one task, one split, one prompting strategy and one fine-tuning framework.")
c.head("2. Objectives")
c.bullets([
    "Adapt three representative VLMs to ten-class defect classification with 4-bit QLoRA.",
    "Compare overall and class-wise performance on a shared held-out test set.",
    "Identify shared and model-specific directional misclassification patterns.",
    "Examine deployment constraints across a consumer workstation and an HPC platform.",
    "Assess whether a smaller VLM offers a competitive accuracy-resource trade-off."])
c.head("3. Dataset")
c.text("The image-level experimental dataset was derived from the local GC10-DET release and contains 2,312 grayscale "
       "steel-surface images in ten defect categories; images were organised into mutually exclusive class folders, giving "
       "a single-label classification task. A stratified split (seed 42) produced 1,845 training, 227 validation and 240 "
       "test images (approx. 8:1:1); the 240-image test set is identical for all models.")
c.pic(FIGS / "fig1_class_distribution.png", COLW,
      cap="Figure 2. Class distribution of the experimental dataset and the held-out test set. Silk spot accounts for "
          "28.2% of the dataset while rolled pit accounts for 1.3%; macro-averaged metrics are therefore reported "
          "alongside accuracy.")
c.head("4. Example Defect Images")
c.pic(FIGS / "fig8_examples.png", COLW,
      cap="Figure 3. Preprocessed GC10-DET examples (CLAHE-enhanced, 512 x 512). Visually similar spot-like defects (water "
          "spot, oil spot, inclusion) form the dominant confusion group.")
c.head("5. Supporting View: Confusion Matrices")
c.pic(FIGS / "fig7_cm_normalized.png", COLW,
      cap="Figure 4. Row-normalised confusion matrices (%): diagonal dominance is comparable across models while "
          "off-diagonal mass concentrates in the two confusion groups of Section 9.")
cols.append(c)
end1 = c.y

# ================= column 2 =================
c = Col(X2, YBODY)
c.head("6. Architecture-Driven Model Selection", before=0)
c.pic(tbl_arch, COLW)
c.head("7. Methodology")
c.text("Image preprocessing. A common reference pipeline converted each image to grayscale, enhanced local contrast with "
       "CLAHE (clip limit 2.0, 8 x 8 tile grid), applied light Gaussian denoising (3 x 3 kernel, sigma 0.8), converted to "
       "three-channel RGB and resized to 512 x 512; each model then applied its native vision processor, so the final "
       "visual-token representation remained model-specific.")
c.text("Expert Prompting (EP). Each model acted as a senior metal-surface inspection expert, received per-class "
       "morphological descriptions and a three-step analysis scaffold, and returned a strict JSON object. Only the "
       "defect_type field was used for quantitative evaluation; auxiliary severity and description fields were excluded "
       "because GC10-DET provides no ground-truth severity annotations.")
c.shade('Prompt (abridged): You are a senior metal-surface quality-inspection expert. Domain knowledge: ten defect classes '
        'with morphology descriptions (e.g. punching_hole: round/oval holes; welding_line: linear weld traces; inclusion: '
        'embedded foreign particles). Steps: (1) inspect for anomalous regions; (2) match the class descriptions; (3) '
        'determine defect type and severity. Output strictly JSON only: {"defect_type": <class>, "severity": <mild|'
        'moderate|severe>, "description": <brief morphology>}.', "F2F2F2", size=12.5)
c.text("Training. All three VLMs were adapted with 4-bit NF4 QLoRA in LLaMA-Factory: LoRA rank 16, alpha 32, dropout 0.05 "
       "on the language-model projections (q/k/v/o/gate/up/down); 3 epochs; LR 1e-4 with cosine schedule and 20-step "
       "warm-up; batch 1 with gradient accumulation 8; cutoff 1024; bf16 with gradient checkpointing; seed 42. Vision "
       "towers were not trained in any model; for InternVL3-8B the vision tower and multimodal projector were additionally "
       "frozen explicitly. Adapters were merged into the base weights before evaluation.")
c.text("Evaluation. Labels were parsed from the JSON defect_type field and scored on the shared test set with accuracy, "
       "Macro-F1, Weighted-F1 and per-class precision, recall and F1.")
c.head("8. Overall Predictive Performance")
c.pic(FIGS / "fig3_overall.png", COLW,
      cap="Figure 5. Overall performance of the three QLoRA-adapted VLMs (vertical axis truncated at 80%). "
          "Per model: Qwen2.5-VL-7B Acc 86.67 / Ma-P 86.19 / Ma-R 83.05 / Ma-F1 84.14 / W-F1 86.63, 32 errors; "
          "Gemma3-4B 87.08 / 85.99 / 84.00 / 83.59 / 86.75, 31 errors; InternVL3-8B 86.67 / 85.57 / 84.44 / 84.80 / "
          "86.52, 32 errors. Macro metrics recomputed from the multiclass confusion matrices; weighted metrics follow "
          "class support.")
c.shade("accuracy is closely matched (86.67-87.08%; gap 0.41 percentage points, i.e. one test image), yet the "
        "accuracy and Macro-F1 rankings disagree, so aggregate accuracy alone cannot rank models on this imbalanced test "
        "set.", "DCE6F1", size=14, bold_prefix="Key result: ")
c.pic(FIGS / "fig4_recall_heatmap.png", COLW - 6.4,
      cap="Figure 6. Class-wise recall. Gemma3-4B led on punching hole, crescent gap and water spot; Qwen2.5-VL-7B reached "
          "perfect recall on silk spot; InternVL3-8B classified all four rolled-pit samples correctly. Inclusion recall "
          "was 69.57% for all three models; rolled pit (n=4) and crease (n=6) should be read cautiously.")
cols.append(c)
end2 = c.y

# ================= column 3 =================
c = Col(X3, YBODY)
c.head("9. Per-Class F1 and Error Analysis", before=0)
c.pic(FIGS / "fig5_f1_heatmap.png", COLW - 7.0,
      cap="Figure 7. Per-class F1-score. No model dominated every class: Qwen2.5-VL-7B led on oil spot, silk spot, crease "
          "and waist folding; Gemma3-4B on punching hole and inclusion; InternVL3-8B on crescent gap, water spot and "
          "rolled pit.", cap_size=12.5)
c.pic(FIGS / "fig6_misclass.png", COLW - 2.0,
      cap="Figure 8. Dominant directional misclassifications per model (true class to predicted class).", cap_size=12.5)
c.text("Two recurrent confusion groups appear: punching hole versus welding line, and inclusion versus spot-like defects. "
       "Both members of each group can produce strong local edge or contrast cues, and inclusions repeatedly predicted as "
       "spots show that small, low-contrast features remain difficult after adaptation.", size=14.5)
c.shade("Entries Cij and Cji are different class-conditional errors and are not "
        "theoretically equal; unequal class support magnifies asymmetry in raw counts, while directional visual ambiguity "
        "sets the dominant error flow.", "FCE4D6", size=13,
        bold_prefix="Why are the confusion matrices asymmetric? ")
c.head("10. Deployment Feasibility")
c.pic(tbl_deploy, COLW)
c.text("Latency is environment-specific and not directly comparable. InternVL3-8B failed to load on the local Windows "
       "workstation (bitsandbytes 4-bit error; BF16 load exceeded the page-file limit) but completed normally on BlueBEAR "
       "A100 - a platform-compatibility issue, not an inherent HPC requirement.", size=14)
c.text("Gemma3-4B: highest accuracy with the smallest parameter count on a consumer GPU.", size=14.5, bold=True)
c.head("11. Excluded Baseline and Prior Research")
c.text("LLaVA-OneVision-7B was planned as a fourth family (SigLIP encoder, projector alignment, AnyRes multi-tile) but "
       "excluded after persistent AnyRes token-alignment failures under the BlueBEAR storage quota - an integration "
       "limitation, not evidence of poor ability. Expert Prompting previously improved VLM defect classification on the "
       "six-class NEU benchmark (+12.7% accuracy for Qwen2.5-VL-7B) [8]; our prompt design was informed by it. LNG-tank "
       "research independently evaluated Qwen-, Gemma-3- and LLaVA-family models under resource-constrained QLoRA [9]; "
       "direct numerical comparison with our setting is inappropriate.", size=14)
c.head("12. Key Findings and Conclusions")
c.text("Headline numbers: 87.08% highest accuracy (Gemma3-4B); 84.80% highest Macro-F1 (InternVL3-8B); 86.19% highest "
       "Macro precision (Qwen2.5-VL-7B); 0.41 pp maximum accuracy gap (one test image); 69.57% inclusion recall, all "
       "three models.", size=14, bold=True)
c.bullets([
    "Three architecture-distinct VLMs matched closely overall under a shared EP and 4-bit QLoRA framework.",
    "Gemma3-4B: highest accuracy (87.08%) and Weighted-F1 (86.75%) with fewest parameters; InternVL3-8B: highest Macro-F1 (84.80%) and Macro recall (84.44%).",
    "Similar accuracy concealed large class-specific differences; recall, F1 and directional errors must be read jointly.",
    "Inclusion is a shared weakness; punching hole and welding line form the main reciprocal confusion group; deployment observations are platform-specific."], size=14)
c.head("13. Limitations and Reproducibility")
c.bullets([
    "Test set limited to 240 images; rolled pit (n=4) and crease (n=6) metrics are unstable; the 0.41-pp gap equals one image.",
    "No ablation of CLAHE/denoising or of the EP schema; native processors yield different visual tokens.",
    "No confidence scores, hence no ROC/PR curves or paired tests; latency from heterogeneous platforms.",
    "Code, prompt template, evaluation scripts and supplementary tables will be released on GitHub."],
    size=13.5)
cols.append(c)
end3 = c.y
print("raw column end y (cm):", round(end1, 1), round(end2, 1), round(end3, 1),
      "| target:", round(TARGET, 1))
for cc in cols:
    cc.distribute(TARGET)
print("distributed end y:", round(cols[0].y, 1), round(cols[1].y, 1), round(cols[2].y, 1))

# ================= footer: references strip across all columns =================
strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(1.2), Cm(STRIP_TOP),
                               Cm(57.0), Cm(STRIP_H))
strip.fill.solid()
strip.fill.fore_color.rgb = RGBColor.from_string("F2F2F2")
strip.line.color.rgb = NAVY
strip.line.width = Pt(1.2)
strip.shadow.inherit = False
tf = strip.text_frame
tf.word_wrap = True
tf.vertical_anchor = 1  # top
tf.margin_left = tf.margin_right = Cm(0.4)
tf.margin_top = Cm(0.2)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
p.space_after = Pt(2)
r = p.add_run()
r.text = "References"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = NAVY
r.font.name = "Calibri"

refs = [
    "[1] X. Lv et al., Deep Metallic Surface Defect Detection: The New Benchmark and Detection Network, Sensors 20(6):1562, 2020, doi:10.3390/s20061562.",
    "[2] T. Dettmers et al., QLoRA: Efficient finetuning of quantized LLMs, NeurIPS 36, 2023.",
    "[3] Y. Zheng et al., LLaMA-Factory: Unified efficient fine-tuning of 100+ language models, ACL 2024 System Demos.",
    "[4] S. Bai et al., Qwen2.5-VL technical report, arXiv:2502.13923, 2025.",
    "[5] Gemma Team, Gemma 3 technical report, arXiv:2503.19786, 2025.",
    "[6] J. Zhu et al., InternVL3, arXiv:2504.10479, 2025.",
    "[7] B. Li et al., LLaVA-OneVision: Easy visual task transfer, arXiv:2408.03326, 2024.",
    "[8] A. H. Aroua and F. Abbas, Enhanced prompting for VLMs in industrial surface defect detection, AIIA 2025, doi:10.1109/AIIA68273.2025.11383671.",
    "[9] H. V. Nguyen et al., Resource-efficient adaptation of LVLMs for multimodal defect inspection in LNG tank manufacturing, Int. J. Adv. Manuf. Technol. 144:5135-5152, 2026, doi:10.1007/s00170-026-17818-y.",
]
groups = [refs[0:3], refs[3:6], refs[6:9]]
for gx, grp in zip((X1 + 0.4, X2 + 0.2, X3 + 0.2), groups):
    tb = slide.shapes.add_textbox(Cm(gx), Cm(STRIP_TOP + 1.15), Cm(COLW - 0.8), Cm(STRIP_H - 1.3))
    tfb = tb.text_frame
    tfb.word_wrap = True
    tfb.margin_left = tfb.margin_right = Cm(0.05)
    tfb.margin_top = tfb.margin_bottom = Cm(0.0)
    for k, rtext in enumerate(grp):
        pp = tfb.paragraphs[0] if k == 0 else tfb.add_paragraph()
        pp.space_after = Pt(3)
        rr = pp.add_run()
        rr.text = rtext
        rr.font.size = Pt(12.5)
        rr.font.name = "Calibri"
        rr.font.color.rgb = BODY

prs.save(str(OUT / "Metal_Defect_VLM_Poster_Final.pptx"))
print("saved. column end y (cm):", round(end1, 1), round(end2, 1), round(end3, 1), "| page height:", PH)
